"""
Fetches Nifty + Sensex candle data from Kite for a given date and writes
an Excel file to runtime_data/candles_YYYY-MM-DD.xlsx.
Returns the file path on success, raises on failure.
"""
import os
import redis
from datetime import date as _date
from dotenv import load_dotenv

BASE_DIR   = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CANDLE_DIR = os.path.join(BASE_DIR, "runtime_data")

load_dotenv(os.path.join(BASE_DIR, ".env"))

r = redis.Redis(host="127.0.0.1", port=6379, decode_responses=True)


def generate_candle_excel(target_date: str | None = None) -> str:
    """
    target_date: 'YYYY-MM-DD', defaults to today.
    Returns absolute path of the written .xlsx file.
    """
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from kiteconnect import KiteConnect

    if target_date is None:
        target_date = _date.today().isoformat()

    auth = r.hgetall("auth") or {}
    if auth.get("state") != "VALID" or not auth.get("token"):
        raise RuntimeError("Kite session not VALID — cannot fetch candles")

    kite = KiteConnect(api_key=os.getenv("API_KEY"))
    kite.set_access_token(auth["token"])

    wb = openpyxl.Workbook()
    wb.remove(wb.active)          # remove default empty sheet

    HDR_FILL  = PatternFill("solid", fgColor="1E3A5F")
    HDR_FONT  = Font(bold=True, color="FFFFFF", size=10)
    HDR_ALIGN = Alignment(horizontal="center")
    COLS      = ["Time", "Open", "High", "Low", "Close", "Volume"]

    def _add_sheet(name: str, candles: list):
        ws = wb.create_sheet(name)
        # header row
        for col_idx, col_name in enumerate(COLS, 1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.font  = HDR_FONT
            cell.fill  = HDR_FILL
            cell.alignment = HDR_ALIGN
        # data rows
        for row_idx, c in enumerate(candles, 2):
            ts = c["date"]
            if hasattr(ts, "strftime"):
                ts = ts.strftime("%H:%M")
            ws.cell(row=row_idx, column=1, value=ts)
            ws.cell(row=row_idx, column=2, value=c["open"])
            ws.cell(row=row_idx, column=3, value=c["high"])
            ws.cell(row=row_idx, column=4, value=c["low"])
            ws.cell(row=row_idx, column=5, value=c["close"])
            ws.cell(row=row_idx, column=6, value=c["volume"])
        # column widths
        ws.column_dimensions["A"].width = 10
        for col in ["B", "C", "D", "E"]:
            ws.column_dimensions[col].width = 12
        ws.column_dimensions["F"].width = 14

    for label, token in [("NIFTY", 256265), ("SENSEX", 265)]:
        for interval, sheet_label in [("minute", "1min"), ("5minute", "5min")]:
            candles = kite.historical_data(
                token, target_date, target_date, interval, continuous=False
            )
            _add_sheet(f"{label} {sheet_label}", candles)

    os.makedirs(CANDLE_DIR, exist_ok=True)
    path = os.path.join(CANDLE_DIR, f"candles_{target_date}.xlsx")
    wb.save(path)
    return path


def get_public_url(file_path: str) -> str:
    """Build the public download URL for a candle file."""
    filename  = os.path.basename(file_path)        # candles_YYYY-MM-DD.xlsx
    date_part = filename.replace("candles_", "").replace(".xlsx", "")
    domain    = os.getenv("REPLIT_DOMAINS", os.getenv("REPLIT_DEV_DOMAIN", "localhost"))
    base_domain = domain.split(",")[0].strip()
    return f"https://{base_domain}/smc/core/candle_file/{date_part}"
