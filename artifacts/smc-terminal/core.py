# core.py
import time
import json
import redis
import pytz
from datetime import datetime as _dt
_IST = pytz.timezone("Asia/Kolkata")
from flask import Blueprint, request, jsonify, Response

# ✅ Verified Constants
from infra.constants import (
    DEFAULT_LOT_NIFTY_OPT, 
    DEFAULT_LOT_SENSEX_OPT,
    TRADE_KEY
)
from services.instrument_lookup import find_option, get_expiries
from services.trading_calendar import get_trading_date
from services.risk_service import calculate_exit_metrics
from services.trade_store import get_trades as _get_trades, save_trades as _save_trades
from services.alert_service import add_alert, remove_alert, get_alerts
from services.fii_dii_service import get_fii_dii
from ops.notion_logger import log_trade_to_notion 
from infra import redis_bus as rbus
from ops.system import snapshot

r = redis.Redis(host="127.0.0.1", port=6379, decode_responses=True)
core = Blueprint("core", __name__)

@core.route("/nifty_card")
def nifty_card():
    stats = rbus.get_json("NIFTY_STATS", {"lp": 0, "change": 0, "p_change": 0, "ts": "--"})
    stats["market"] = snapshot()["market"]
    return jsonify(stats)

@core.route("/trade_state")
def trade_state():
    """Handles LTP, MFE/MAE points, filtering (date), and sorting."""

    trades = _get_trades()

    # ✅ FILTER: Only active + today's trades
    today = get_trading_date()
    trades = [
        t for t in trades
        if t["status"] == "ACTIVE" or t.get("date") == today
    ]

    updated = False

    for t in trades:
        ltp_val = r.get(f"ltp:{t['token']}")
        if ltp_val:
            ltp = float(ltp_val)
            t["ltp"] = ltp

            entry = float(t["entryPrice"])
            qty = int(t["lots"]) * int(t.get("multiplier", 1))

            if t["status"] == "ACTIVE":
                t["net_pnl"] = round(
                    (ltp - entry) * qty if t["direction"] == "LONG"
                    else (entry - ltp) * qty,
                    2
                )

            # ✅ Update MFE / MAE
            t["mfe"] = round(
                max(
                    float(t.get("mfe", 0)),
                    max(
                        0,
                        ltp - entry if t["direction"] == "LONG"
                        else entry - ltp
                    )
                ),
                2
            )

            t["mae"] = round(
                min(
                    float(t.get("mae", 0)),
                    min(
                        0,
                        ltp - entry if t["direction"] == "LONG"
                        else entry - ltp
                    )
                ),
                2
            )

            updated = True

    # ✅ Sort: ACTIVE first, then newest
    trades.sort(key=lambda x: (x["status"] != "ACTIVE", -x.get("created_at", 0)))

    if updated:
        _save_trades(trades)

    return jsonify(trades)



@core.route("/trade", methods=["POST"])
def create_trade():
    data = request.json
    instr = find_option(data["index"], data["strike"], data["option_type"], data.get("expiry"))
    if not instr: return jsonify({"error": "Instrument not found"}), 400

    mult = DEFAULT_LOT_NIFTY_OPT if data["index"] == "NIFTY" else DEFAULT_LOT_SENSEX_OPT
    
    # ✅ ADDED: Capture SL and TG from incoming request
    new_trade = {
        "id": f"T{int(time.time()*1000)}", 
        "symbol": instr["symbol"], 
        "token": instr["token"],
        "index": data["index"],
        "entryPrice": float(data["entryPrice"]), 
        "lots": int(data["lots"]), 
        "sl": float(data.get("sl", 0)),  # <-- NEW
        "tg": float(data.get("tg", 0)),  # <-- NEW
        "multiplier": mult,
        "direction": data.get("direction", "LONG"), 
        "status": "ACTIVE", 
        "created_at": time.time(), 
        "setup": data.get("setup", "SMC"), 
        "entry_emotion": data.get("entry_emotion", "Calm"), 
        "net_pnl": 0, "mfe": 0, "mae": 0, "ltp": 0,
        "entry_time": data.get("entry_time", _dt.now(_IST).strftime("%H:%M:%S")),
        "date": get_trading_date()
    }
    
    trades = _get_trades()
    trades.append(new_trade)
    _save_trades(trades)
    return jsonify(new_trade)

@core.route("/trade/exit", methods=["POST"])
def exit_trade():
    data = request.json
    trades = _get_trades()
    updated = False
    for t in trades:
        if t["id"] == data["id"] and t["status"] == "ACTIVE":
            metrics = calculate_exit_metrics(t, float(data["exit_price"]))
            t.update({
                "status": "CLOSED", 
                "exit_price": float(data["exit_price"]), 
                "exit_time": data.get("exit_time", _dt.now(_IST).strftime("%H:%M:%S")),
                "exit_emotion": data.get("exit_emotion", "Disciplined"), 
                "net_pnl": metrics["net_pnl"],
                "brokerage": metrics["brokerage"]
            })
            
            # ✅ PASS DATA TO NOTION
            log_trade_to_notion(t) 
            updated = True; break
            
    if updated: _save_trades(trades)
    return jsonify({"status": "ok"})

@core.route("/alerts/get")
def get_alerts_route():
    return jsonify(get_alerts())

@core.route("/alerts/save", methods=["POST"])
def save_alert():
    add_alert(request.json)
    return jsonify({"status": "ok"})

@core.route("/alerts/delete", methods=["POST"])
def delete_alert():
    remove_alert(request.json["id"])
    return jsonify({"status": "ok"})

@core.route("/fii_dii")
def fii_dii():
    return jsonify(get_fii_dii())

@core.route("/valid_expiries")
def valid_expiries():
    options = get_expiries(request.args.get("index"))
    return jsonify({"options": options})

@core.route("/trade/edit", methods=["POST"])
def edit_trade():
    data = request.json
    trades = _get_trades()
    updated = False
    for t in trades:
        if t["id"] == data["id"]:
            for field in ["sl", "tg", "entryPrice", "lots", "setup", "entry_emotion", "entry_time"]:
                if field in data:
                    t[field] = float(data[field]) if field in ("sl", "tg", "entryPrice") else \
                               int(data[field])   if field == "lots" else data[field]
            updated = True
            break
    if updated:
        _save_trades(trades)
    return jsonify({"status": "ok"})

@core.route("/last_tick")
def last_tick():
    nifty_stats  = rbus.get_json("NIFTY_STATS",  None)
    sensex_stats = rbus.get_json("SENSEX_STATS", None)

    # Fall back to raw LTP keys when STATS haven't been written yet
    if not nifty_stats or not nifty_stats.get("lp"):
        raw = rbus.get("ltp:256265")
        nifty_stats = {"lp": float(raw)} if raw else {"lp": 0}

    if not sensex_stats or not sensex_stats.get("lp"):
        raw = rbus.get("ltp:265")
        sensex_stats = {"lp": float(raw)} if raw else {"lp": 0}

    return jsonify({"nifty": nifty_stats, "sensex": sensex_stats})

@core.route("/expiries")
def expiries():
    return jsonify(get_expiries(request.args.get("index")))

@core.route("/system")
def system_health():
    return jsonify(snapshot())


@core.route("/export_trades")
def export_trades():
    """
    Download trades as Excel.
    Query params: from_date (YYYY-MM-DD), to_date (YYYY-MM-DD).
    Defaults to all trades if no range given.
    """
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, numbers
    import io

    from_date = request.args.get("from_date", "")
    to_date   = request.args.get("to_date", "")

    all_trades = _get_trades()

    # Filter by date range
    filtered = []
    for t in all_trades:
        d = t.get("date", "")
        if from_date and d < from_date:
            continue
        if to_date and d > to_date:
            continue
        filtered.append(t)

    # Sort oldest → newest
    filtered.sort(key=lambda x: (x.get("date", ""), x.get("entry_time", "")))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Trades"

    HDR_FILL  = PatternFill("solid", fgColor="1E3A5F")
    HDR_FONT  = Font(bold=True, color="FFFFFF", size=10)
    WIN_FILL  = PatternFill("solid", fgColor="D4EDDA")
    LOSS_FILL = PatternFill("solid", fgColor="F8D7DA")
    ACTV_FILL = PatternFill("solid", fgColor="FFF3CD")

    HEADERS = [
        "Date", "Symbol", "Direction", "Setup",
        "Entry Time", "Entry Price", "SL", "TG",
        "Lots", "Exit Time", "Exit Price", "Net P&L",
        "MFE (pts)", "MAE (pts)", "Brokerage",
        "Entry Emotion", "Exit Emotion", "Status"
    ]

    for col_idx, h in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font      = HDR_FONT
        cell.fill      = HDR_FILL
        cell.alignment = Alignment(horizontal="center")

    for row_idx, t in enumerate(filtered, 2):
        row = [
            t.get("date", ""),
            t.get("symbol", ""),
            t.get("direction", ""),
            t.get("setup", ""),
            t.get("entry_time", ""),
            t.get("entryPrice", ""),
            t.get("sl", ""),
            t.get("tg", ""),
            t.get("lots", ""),
            t.get("exit_time", ""),
            t.get("exit_price", ""),
            t.get("net_pnl", ""),
            t.get("mfe", ""),
            t.get("mae", ""),
            t.get("brokerage", ""),
            t.get("entry_emotion", ""),
            t.get("exit_emotion", ""),
            t.get("status", ""),
        ]
        for col_idx, val in enumerate(row, 1):
            ws.cell(row=row_idx, column=col_idx, value=val)

        # Colour-code rows
        status = t.get("status", "")
        pnl    = t.get("net_pnl", 0) or 0
        if status == "ACTIVE":
            fill = ACTV_FILL
        elif pnl >= 0:
            fill = WIN_FILL
        else:
            fill = LOSS_FILL
        for col_idx in range(1, len(HEADERS) + 1):
            ws.cell(row=row_idx, column=col_idx).fill = fill

    # Column widths
    widths = [12, 26, 10, 12, 11, 13, 10, 10, 6,
              11, 13, 12, 11, 11, 12, 16, 16, 10]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

    # Summary row
    if filtered:
        summary_row = len(filtered) + 3
        ws.cell(row=summary_row, column=1, value="SUMMARY").font = Font(bold=True)
        ws.cell(row=summary_row, column=2, value=f"{len(filtered)} trades")
        closed = [t for t in filtered if t.get("status") == "CLOSED"]
        wins   = [t for t in closed if (t.get("net_pnl") or 0) > 0]
        total_pnl = sum(t.get("net_pnl", 0) or 0 for t in closed)
        ws.cell(row=summary_row + 1, column=1, value="Closed")
        ws.cell(row=summary_row + 1, column=2, value=len(closed))
        ws.cell(row=summary_row + 2, column=1, value="Win Rate")
        ws.cell(row=summary_row + 2, column=2,
                value=f"{round(len(wins)/len(closed)*100)}%" if closed else "—")
        ws.cell(row=summary_row + 3, column=1, value="Net P&L")
        pnl_cell = ws.cell(row=summary_row + 3, column=2, value=round(total_pnl, 2))
        pnl_cell.font = Font(bold=True,
                             color="006400" if total_pnl >= 0 else "8B0000")

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    label = f"{from_date}_to_{to_date}" if from_date or to_date else "all"
    filename = f"trades_{label}.xlsx"
    return Response(
        buf.read(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )


@core.route("/candle_file/<date_str>")
def candle_file(date_str):
    """Serve a pre-generated candle Excel file for the given date."""
    import os, re
    from flask import send_file
    if not re.match(r"^\d{4}-\d{2}-\d{2}$", date_str):
        return jsonify({"error": "invalid date"}), 400
    path = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                        "runtime_data", f"candles_{date_str}.xlsx")
    if not os.path.exists(path):
        return jsonify({"error": f"No candle file for {date_str} yet — "
                                  "it is generated at 3:45 PM IST after market close."}), 404
    return send_file(path, as_attachment=True,
                     download_name=f"kite_candles_{date_str}.xlsx")

